from openpyxl import Workbook


def export_to_excel(rows, headers, output_path):
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(output_path)


_ENCABEZADOS_DETALLE = (
    "Nombre", "Cédula", "Empresa Convenio", "Microseguro", "Motivo No Aplica", "Fecha",
)


def exportar_reporte_mensual(resultado, nombre_mes, anio, agente, output_path):
    """Exporta el Reporte Mensual de Casos (ver db/reporte_mensual.py) a un
    .xlsx con una hoja "Resumen" y una hoja por categoría. Pedido explícito
    del usuario: este archivo queda como referencia fija en su computadora,
    independiente de lo que la base de datos muestre después — puede abrir
    el de un mes junto al de otro y compararlos él mismo, sin depender de
    que nada en la app cambie por debajo."""
    wb = Workbook()

    resumen = wb.active
    resumen.title = "Resumen"
    resumen.append(["Reporte Mensual de Casos"])
    resumen.append([f"{nombre_mes} {anio}", f"Agente: {agente}"])
    resumen.append([])
    resumen.append(["Categoría", "Cantidad"])
    resumen.append(["Desembolsados", len(resultado["desembolsados"])])
    resumen.append(["  Con microseguro", resultado["con_microseguro"]])
    resumen.append(["  Sin microseguro", resultado["sin_microseguro"]])
    resumen.append(["No aplica", len(resultado["no_aplica"])])
    resumen.append(["Cliente desistió", len(resultado["cliente_desistio"])])
    resumen.append(["Pendientes (no depende del mes elegido)", len(resultado["pendientes"])])

    _agregar_hoja_detalle(wb, "Desembolsados", resultado["desembolsados"])
    _agregar_hoja_detalle(wb, "No aplica", resultado["no_aplica"])
    _agregar_hoja_detalle(wb, "Cliente desistio", resultado["cliente_desistio"])
    _agregar_hoja_detalle(wb, "Pendientes", resultado["pendientes"])

    wb.save(output_path)


def _agregar_hoja_detalle(wb, titulo, entradas):
    hoja = wb.create_sheet(titulo)
    hoja.append(list(_ENCABEZADOS_DETALLE))
    for entrada in entradas:
        hoja.append([
            entrada["nombre"], entrada["cedula"], entrada["empresa_convenio"],
            entrada["microseguro"], entrada["motivo_no_aplica"] or "", entrada["fecha"] or "",
        ])
