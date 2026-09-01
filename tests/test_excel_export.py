import openpyxl

from gestor_credito.export.excel_export import exportar_reporte_mensual


def _resultado_de_ejemplo():
    return {
        "total_registrados": 3,
        "desembolsados": [
            {
                "clave_caso": "P-1", "nombre": "Juan Perez", "cedula": "001",
                "empresa_convenio": "MIDESA", "microseguro": "Sí",
                "motivo_no_aplica": None, "fecha": "2026-08-20",
            },
        ],
        "con_microseguro": 1,
        "sin_microseguro": 0,
        "no_aplica": [
            {
                "clave_caso": "P-2", "nombre": "Ana Lopez", "cedula": "002",
                "empresa_convenio": "MIDESA", "microseguro": "No",
                "motivo_no_aplica": "Ingresos insuficientes", "fecha": "2026-08-05",
            },
        ],
        "cliente_desistio": [],
        "pendientes": [
            {
                "clave_caso": "P-3", "nombre": "Beto Cruz", "cedula": "003",
                "empresa_convenio": "MIDESA", "microseguro": "No",
                "motivo_no_aplica": None, "fecha": None,
            },
        ],
    }


def test_exportar_reporte_mensual_crea_hoja_resumen_y_detalle(tmp_path):
    ruta = tmp_path / "reporte.xlsx"
    exportar_reporte_mensual(_resultado_de_ejemplo(), "Agosto", 2026, "Maria Gomez", str(ruta))

    wb = openpyxl.load_workbook(ruta)
    assert wb.sheetnames == ["Resumen", "Desembolsados", "No aplica", "Cliente desistio", "Pendientes"]

    resumen = wb["Resumen"]
    valores_resumen = [fila[0].value for fila in resumen.iter_rows()]
    assert "Reporte Mensual de Casos" in valores_resumen
    assert "Total de casos registrados en el mes" in valores_resumen

    # Encabezados: Caso, Nombre, Cédula, Empresa Convenio, Microseguro, Motivo No Aplica, Fecha.
    desembolsados = wb["Desembolsados"]
    assert desembolsados["A1"].value == "Caso"
    assert desembolsados["B1"].value == "Nombre"
    assert desembolsados["A2"].value == "P-1"
    assert desembolsados["B2"].value == "Juan Perez"
    assert desembolsados["G2"].value == "2026-08-20"

    pendientes = wb["Pendientes"]
    assert pendientes["A2"].value == "P-3"
    assert pendientes["B2"].value == "Beto Cruz"
    assert pendientes["G2"].value is None
